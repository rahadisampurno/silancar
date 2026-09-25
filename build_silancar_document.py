from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUTPUT = Path("deliverables/Silancar_Hasil_Interview_dan_Rancangan_Solusi.docx")

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "666666"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if value and keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    elif not value and keep is not None:
        p_pr.remove(keep)


def set_repeat_table_rows(table):
    if table.rows:
        set_repeat_table_header(table.rows[0])


def set_font(run, name="Arial", size=None, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", style=None, bold_lead=None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    if keep:
        set_keep_with_next(p)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_font(r)
    return p


_number_counter = 0


def reset_numbering():
    global _number_counter
    _number_counter = 0


def add_number(doc, text, level=0):
    global _number_counter
    _number_counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.24 if level == 0 else 0.48)
    p.paragraph_format.first_line_indent = Inches(-0.24)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(f"{_number_counter}.  {text}")
    set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(str(header))
        set_font(r, size=font_size, bold=True, color=WHITE)
    prevent_row_split(hdr)
    for row_idx, row_data in enumerate(rows):
        added_row = table.add_row()
        prevent_row_split(added_row)
        cells = added_row.cells
        for i, value in enumerate(row_data):
            cell = cells[i]
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            p.alignment = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_font(r, size=font_size)
    set_repeat_table_rows(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_status_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + "  ")
    set_font(r, size=9.5, bold=True, color=MID_GRAY)
    r2 = p.add_run(value)
    set_font(r2, size=9.5)
    return p


def set_document_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)
    title_ppr = title._element.get_or_add_pPr()
    title_pbdr = OxmlElement("w:pBdr")
    title_bottom = OxmlElement("w:bottom")
    title_bottom.set(qn("w:val"), "nil")
    title_pbdr.append(title_bottom)
    title_ppr.append(title_pbdr)

    for name, size, before, after in (("Heading 1", 16, 16, 7), ("Heading 2", 12.5, 12, 5), ("Heading 3", 11, 9, 3)):
        s = styles[name]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(BLACK)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        s = styles[list_style]
        s.font.name = "Arial"
        s.font.size = Pt(10.5)
        s.paragraph_format.space_after = Pt(3)
        s.paragraph_format.line_spacing = 1.1

    if "Small Note" not in styles:
        s = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        s.font.size = Pt(9)
        s.font.color.rgb = RGBColor.from_string(MID_GRAY)
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.line_spacing = 1.05


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
set_document_styles(doc)

# Cover page
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(52)
r = p.add_run("SILANCAR")
set_font(r, size=12, bold=True, color=BLUE)

title = doc.add_paragraph(style="Title")
title.add_run("Hasil Interview dan Rancangan Awal Solusi")
title_ppr = title._p.get_or_add_pPr()
title_pbdr = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_pbdr.append(title_bottom)
title_ppr.append(title_pbdr)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(26)
r = subtitle.add_run("Sistem omnichannel, pengelolaan penjualan, operasional, produksi, KPI, dan OKR")
set_font(r, size=14, color=MID_GRAY)

add_status_line(doc, "Pembaca", "Owner Silancar")
add_status_line(doc, "Jenis dokumen", "Hasil interview dan rekomendasi solusi awal")
add_status_line(doc, "Tanggal interview", "Belum dikonfirmasi")
add_status_line(doc, "Status", "Dokumen pembahasan awal untuk validasi kebutuhan")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(34)
r = p.add_run("Tujuan dokumen")
set_font(r, size=11, bold=True)
add_paragraph(doc, "Dokumen ini menerjemahkan hasil interview menjadi rancangan yang dapat diperiksa oleh owner. Isinya menjelaskan masalah, ruang lingkup sistem, alur kerja, pembagian peran, ukuran kinerja, dan keputusan yang masih perlu dikonfirmasi sebelum pengembangan dimulai.")

doc.add_page_break()

add_heading(doc, "Ringkasan Eksekutif", 1)
add_paragraph(doc, "Silancar membutuhkan satu sistem kerja yang menyatukan percakapan pelanggan dari WhatsApp, Instagram, dan Facebook dengan proses penjualan, pemenuhan pesanan, produksi, dan evaluasi kinerja tim. Saat ini data sudah dicatat melalui beberapa spreadsheet untuk leads, pipeline, target, dan bonus, tetapi informasi masih tersebar sehingga owner dan admin sulit melihat kondisi bisnis secara utuh dan konsisten.")
add_paragraph(doc, "Rekomendasi awal adalah membangun dashboard operasional dengan kotak masuk omnichannel, database pelanggan dan leads, pipeline penjualan, manajemen order dan produksi, kontrol kualitas PIC brand, serta dashboard KPI dan OKR. Sistem harus menyimpan pemilik pekerjaan, status, tenggat, riwayat aktivitas, dan sumber data agar setiap angka dapat ditelusuri.")
add_paragraph(doc, "Versi pertama sebaiknya memprioritaskan kendali leads dan follow-up, karena volume masuk disebut sekitar 30 sampai 100 leads per hari. Modul produksi dan KPI kemudian dihubungkan ke data transaksi yang sudah stabil. Target angka, rumus bonus, struktur tim final, serta definisi zero mistakes masih perlu diputuskan oleh owner.")

add_heading(doc, "Cara Membaca Dokumen", 2)
add_table(doc,
          ["Penanda", "Arti"],
          [
              ("Temuan", "Informasi yang berasal dari catatan interview atau tampilan proses saat ini."),
              ("Interpretasi", "Makna kerja yang ditarik dari temuan dan perlu dikonfirmasi bila sumbernya ambigu."),
              ("Rekomendasi", "Rancangan awal yang diusulkan untuk menyelesaikan kebutuhan."),
              ("Belum dikonfirmasi", "Keputusan atau nilai yang belum tersedia dan tidak boleh dianggap final."),
          ], [1.35, 5.45], font_size=9.5)

add_heading(doc, "Hasil Interview", 1)
add_heading(doc, "Profil Operasi", 2)
add_bullet(doc, "Usaha bergerak pada trading garment, penjualan kain, dan kerja sama produksi dengan konveksi.")
add_bullet(doc, "Fungsi kerja yang disebut mencakup sales, konten atau digital marketing, operasional, PIC brand, produksi, dan admin.")
add_bullet(doc, "Sales dibedakan menjadi sales online dan sales canvassing. Digital marketing berperan menghasilkan atau mendukung leads.")
add_bullet(doc, "Sales online dan sales offline atau canvassing perlu dapat berkolaborasi ketika dibutuhkan untuk mencapai closing.")
add_bullet(doc, "PIC brand diharapkan menangani brand yang masuk tanpa kesalahan. Definisi dan SOP penilaian 100 persen belum tersedia.")

add_heading(doc, "Proses dan Data Saat Ini", 2)
add_paragraph(doc, "Screenshot menunjukkan bahwa tim telah menggunakan spreadsheet terpisah untuk mencatat detail leads, rekap leads mingguan per kanal, pipeline penjualan, target mingguan, dan skema bonus. Struktur ini membuktikan kebutuhan pelaporan sudah ada, tetapi data masih harus dipindahkan atau dirangkum secara manual.")
add_table(doc,
          ["Area", "Data yang terlihat", "Risiko proses"],
          [
              ("Leads", "Tanggal, sumber, nama, kontak, kebutuhan, tindakan, kanal, layanan, volume, status", "Duplikasi, format tidak seragam, follow-up terlewat"),
              ("Rekap kanal", "Jumlah leads WhatsApp, Instagram, Messenger, total, hot atau warm lead", "Rekap terlambat dan definisi antar-tim dapat berbeda"),
              ("Pipeline", "Brand, sales, PO, kuantitas, nilai, status, tanggal kirim, keterangan", "Status penjualan dan produksi belum terhubung"),
              ("Target", "Target bulanan dan mingguan per orang atau fungsi", "Pencapaian sulit diaudit bila sumber transaksi terpisah"),
              ("Insentif", "Rentang pencapaian, bonus, dan persentase", "Rumus dapat ambigu atau salah bila dihitung manual"),
          ], [1.1, 3.25, 2.45], font_size=8.9)

add_heading(doc, "Masalah Utama", 2)
reset_numbering()
add_number(doc, "Percakapan pelanggan tersebar di tiga kanal sehingga pemantauan dan pembagian follow-up tidak terpusat.")
add_number(doc, "Leads, aktivitas sales, nilai pipeline, pesanan, dan produksi belum membentuk satu riwayat yang tersambung.")
add_number(doc, "Kolaborasi sales online dan canvassing belum memiliki aturan handover, pembagian kontribusi, dan kepemilikan closing yang baku.")
add_number(doc, "Owner belum memiliki dashboard ringkas untuk membandingkan target, realisasi, kendala, dan beban kerja per fungsi.")
add_number(doc, "KPI per jabatan, formula penilaian, serta definisi keberhasilan PIC brand dan produksi belum dibakukan.")

add_heading(doc, "Interpretasi yang Perlu Divalidasi", 2)
add_table(doc,
          ["Catatan awal", "Interpretasi dalam rancangan", "Status"],
          [
              ("Admin tidak dapat memantau semua chat", "Kebutuhan dipahami sebagai admin perlu dapat memantau seluruh chat sesuai hak akses.", "Perlu konfirmasi"),
              ("Tim ada 2 sales sama Dinar konten sama operasional", "Ada dua personel sales serta fungsi konten dan operasional, tetapi nama dan struktur final belum jelas.", "Perlu konfirmasi"),
              ("Brand yang datang 100 persen zero mistakes", "Targetnya adalah pesanan diproses sesuai requirement tanpa kesalahan yang disebabkan proses internal.", "Perlu definisi"),
              ("30 sampai 100 leads", "Angka diperlakukan sebagai kisaran leads masuk per hari, bukan target final.", "Perlu validasi"),
          ], [2.15, 3.55, 1.1], font_size=8.9,
          alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])

add_heading(doc, "Tujuan dan Batas Solusi", 1)
add_heading(doc, "Tujuan Bisnis", 2)
add_bullet(doc, "Tidak ada leads yang hilang atau tidak memiliki tindak lanjut.")
add_bullet(doc, "Owner dapat melihat kondisi sales, order, produksi, dan kinerja tim dari satu dashboard.")
add_bullet(doc, "Setiap pekerjaan memiliki penanggung jawab, target waktu, status, dan bukti penyelesaian.")
add_bullet(doc, "Sales online dan canvassing dapat berkolaborasi tanpa kehilangan riwayat komunikasi atau kejelasan kontribusi.")
add_bullet(doc, "Penilaian KPI dan OKR dihitung dari data operasional yang dapat ditelusuri.")

add_heading(doc, "Cakupan Awal", 2)
add_table(doc,
          ["Termasuk", "Belum termasuk sampai disetujui"],
          [
              ("Kotak masuk WhatsApp, Instagram, dan Facebook dalam satu tampilan", "Akuntansi lengkap dan pelaporan pajak"),
              ("Leads, pelanggan, follow-up, penugasan, dan pipeline penjualan", "Payroll dan sistem HR lengkap"),
              ("Order, requirement brand, sample, produksi, QC, dan pengiriman", "Marketplace dan e-commerce yang belum disebut"),
              ("Target, KPI, OKR, notifikasi, laporan, dan audit aktivitas", "Otomasi kampanye iklan dan content publishing lengkap"),
          ], [3.4, 3.4], font_size=9.1)

add_heading(doc, "Pengguna dan Hak Akses", 2)
add_table(doc,
          ["Peran", "Akses utama", "Batasan yang disarankan"],
          [
              ("Owner", "Seluruh dashboard, target, KPI, laporan, dan audit", "Perubahan formula KPI melalui persetujuan"),
              ("Admin", "Semua inbox, leads, distribusi chat, data pelanggan, dan laporan operasional", "Tidak mengubah target atau formula tanpa izin"),
              ("Digital marketing", "Sumber kampanye, leads yang dihasilkan, kualitas leads, dan konten terkait", "Tidak mengakses biaya atau data sensitif di luar kebutuhan"),
              ("Sales online", "Inbox dan leads yang ditugaskan, pipeline, quotation, follow-up", "Akses percakapan mengikuti tim dan kebijakan owner"),
              ("Sales canvassing", "Prospek lapangan, kunjungan, follow-up, dan peluang kolaborasi", "Tidak mengubah catatan sales lain tanpa handover"),
              ("PIC brand", "Requirement, sample, persetujuan, perubahan, dan koordinasi order", "Tidak menutup order sebelum checklist lengkap"),
              ("Produksi", "Jadwal, kuantitas, progres, masalah, QC, dan hasil", "Tidak mengubah nilai penjualan"),
          ], [1.5, 3.35, 1.95], font_size=8.35)

add_heading(doc, "Rancangan Modul Dashboard", 1)
modules = [
    ("Inbox Omnichannel", "Menampilkan percakapan WhatsApp, Instagram, dan Facebook; pencarian; filter; label; assignment; template balasan; catatan internal; dan SLA respons."),
    ("Leads dan Pelanggan", "Menyimpan profil, sumber, kebutuhan, volume, riwayat komunikasi, pemilik lead, tingkat minat, tindak lanjut berikutnya, dan potensi duplikasi."),
    ("Tugas dan Follow Up", "Membuat jadwal follow-up, pengingat, daftar kerja harian, overdue, eskalasi, dan handover antar-sales."),
    ("Pipeline Penjualan", "Mengelola tahap peluang, estimasi nilai dan kuantitas, quotation, sample, probabilitas, alasan kalah, dan forecast."),
    ("Order dan Produksi", "Menghubungkan deal dengan PO, requirement, material, vendor konveksi, jadwal, progres, QC, pengiriman, dan penyelesaian."),
    ("Kontrol PIC Brand", "Menyediakan checklist requirement, catatan perubahan, approval sample, bukti persetujuan, versioning, dan pencatatan kesalahan."),
    ("KPI dan OKR", "Menghitung target dan realisasi per peran, menampilkan bobot, skor, tren, sumber data, dan catatan evaluasi."),
    ("Dashboard Owner", "Ringkasan leads, SLA, conversion, pipeline, revenue, order berisiko, produksi terlambat, defect, target, dan kinerja tim."),
    ("Master Data dan Audit", "Mengatur user, role, kanal, status, jenis layanan, sumber lead, formula, notifikasi, serta riwayat perubahan."),
]
add_table(doc, ["Modul", "Fungsi utama"], modules, [1.7, 5.1], font_size=9.0)

add_heading(doc, "Dashboard Prioritas", 2)
add_table(doc,
          ["Dashboard", "Indikator yang ditampilkan", "Keputusan yang didukung"],
          [
              ("Owner", "Leads hari ini, response SLA, overdue, conversion, nilai pipeline, order aktif, keterlambatan, KPI", "Prioritas, kapasitas, target, dan eskalasi"),
              ("Sales", "Inbox ditugaskan, follow-up hari ini, aging lead, hot leads, target, forecast", "Kontak berikutnya dan fokus closing"),
              ("Operasional", "Order baru, requirement belum lengkap, sample, jadwal produksi, hambatan, QC, kirim", "Urutan kerja dan penyelesaian kendala"),
              ("Digital marketing", "Leads per sumber, qualified rate, conversion per kampanye, tren kebutuhan", "Evaluasi kualitas sumber leads"),
          ], [1.25, 3.55, 2.0], font_size=8.8)

add_heading(doc, "Alur Kerja Utama", 1)
add_heading(doc, "Alur Leads sampai Closing", 2)
lead_steps = [
    "Pesan masuk dari WhatsApp, Instagram, atau Facebook dan otomatis dibuat sebagai percakapan serta calon lead.",
    "Sistem mencari kecocokan nomor atau akun sosial untuk mencegah duplikasi pelanggan.",
    "Admin atau aturan distribusi menetapkan pemilik lead berdasarkan kanal, beban, wilayah, atau jenis kebutuhan.",
    "Sales merespons, melengkapi kebutuhan, volume, anggaran, tenggat, dan klasifikasi lead.",
    "Sales menetapkan tindakan berikutnya dan tanggal follow-up. Lead tanpa aktivitas melewati SLA masuk ke daftar overdue.",
    "Jika perlu bantuan, sales membuat permintaan kolaborasi kepada sales lain tanpa memindahkan atau menghapus riwayat.",
    "Lead yang memenuhi kriteria dibuat menjadi opportunity dan bergerak melalui quotation, sample, negosiasi, lalu won atau lost.",
    "Deal won menghasilkan order. Deal lost wajib memiliki alasan agar owner dapat membaca pola kegagalan.",
]
reset_numbering()
for step in lead_steps:
    add_number(doc, step)

add_heading(doc, "Status Pipeline yang Disarankan", 3)
add_paragraph(doc, "New > Assigned > Contacted > Qualified > Warm atau Hot > Quotation > Sample > Negotiation > Won atau Lost")
add_paragraph(doc, "Status seperti only read sebaiknya tidak dipakai sebagai tahap pipeline karena tidak menjelaskan kondisi bisnis. Jika masih dibutuhkan, gunakan sebagai label aktivitas, misalnya pesan sudah dibaca tetapi belum dijawab.", style="Small Note")

add_heading(doc, "Alur Kolaborasi Sales", 2)
reset_numbering()
add_number(doc, "Sales pemilik lead mengajukan kolaborasi dan memilih alasan, misalnya bantuan kunjungan, relasi, negosiasi, atau closing.")
add_number(doc, "Sales pendamping menerima tugas dengan ruang lingkup dan tenggat yang jelas.")
add_number(doc, "Semua aktivitas dicatat pada lead yang sama. Sistem tidak membuat pipeline kedua untuk peluang yang sama.")
add_number(doc, "Sebelum closing, sistem meminta konfirmasi peran primary owner, contributor, dan aturan pembagian insentif.")
add_number(doc, "Owner atau atasan menyelesaikan konflik kontribusi berdasarkan riwayat aktivitas yang tercatat.")

add_heading(doc, "Alur Order Trading dan Produksi", 2)
add_table(doc,
          ["Tahap", "Kontrol wajib", "Hasil"],
          [
              ("Order diterima", "PO, item, bahan, warna, ukuran, jumlah, harga, tenggat, alamat", "Order tervalidasi"),
              ("Pilih jalur", "Trading stok atau produksi konveksi", "Rute proses ditentukan"),
              ("Sample", "Referensi, spesifikasi, revisi, approval pelanggan", "Sample disetujui"),
              ("Material", "Ketersediaan kain dan aksesori, vendor, estimasi datang", "Material siap"),
              ("Produksi", "Jadwal, target harian, output, hambatan, PIC", "Barang selesai diproduksi"),
              ("Quality control", "Kuantitas, ukuran, warna, jahitan, branding, kemasan", "Lulus, rework, atau reject"),
              ("Pengiriman", "Alamat, jumlah, dokumen, kurir, tanggal, bukti", "Barang terkirim"),
              ("Penutupan", "Penerimaan, komplain, sisa kewajiban, evaluasi", "Order selesai"),
          ], [1.25, 3.7, 1.85], font_size=8.8)

add_heading(doc, "Kontrol Zero Mistakes untuk PIC Brand", 2)
add_paragraph(doc, "Zero mistakes perlu didefinisikan sebagai order yang memenuhi requirement pelanggan tanpa kesalahan internal yang memicu revisi, rework, retur, atau keterlambatan. Kesalahan dari perubahan pelanggan setelah approval harus dicatat terpisah agar penilaian adil.")
add_bullet(doc, "Checklist requirement wajib lengkap sebelum quotation atau produksi dilanjutkan.")
add_bullet(doc, "Setiap revisi memiliki versi, waktu, pengaju, penerima, dan bukti persetujuan.")
add_bullet(doc, "Sample atau mockup harus disetujui sebelum produksi massal.")
add_bullet(doc, "Perubahan setelah approval harus menghasilkan change request dan dampak biaya atau jadwal.")
add_bullet(doc, "Error log membedakan internal error, vendor error, dan customer change.")

add_heading(doc, "Aturan Bisnis Utama", 1)
rules = [
    ("Kepemilikan", "Setiap lead, opportunity, order, dan tugas wajib memiliki satu primary owner."),
    ("Follow-up", "Lead aktif wajib memiliki next action dan tanggal follow-up."),
    ("Duplikasi", "Nomor telepon atau akun sosial yang sama memicu pemeriksaan sebelum membuat kontak baru."),
    ("Handover", "Perpindahan owner wajib menyimpan alasan, waktu, pemberi, dan penerima."),
    ("Kolaborasi", "Kontributor dapat bekerja pada lead yang sama tanpa mengubah primary owner."),
    ("Perubahan order", "Perubahan setelah approval wajib dicatat sebagai versi baru dan disetujui ulang."),
    ("Penutupan", "Lead lost atau cancelled wajib memiliki alasan; order tidak dapat selesai sebelum checklist penutupan terpenuhi."),
    ("KPI", "Setiap nilai KPI harus menunjukkan periode, formula, target, realisasi, bobot, dan sumber data."),
    ("Audit", "Perubahan status, owner, target, formula, dan nilai transaksi disimpan dalam audit log."),
]
add_table(doc, ["Aturan", "Definisi"], rules, [1.35, 5.45], font_size=9.1)

add_heading(doc, "Rancangan KPI Awal", 1)
add_paragraph(doc, "Tabel berikut adalah rancangan untuk diskusi. Target numerik tidak ditetapkan karena baseline, jam kerja, kapasitas, margin, dan prioritas owner belum dikonfirmasi. Setelah sistem berjalan, gunakan data dua sampai empat minggu untuk menetapkan baseline yang realistis.")

add_heading(doc, "Prinsip Penilaian", 2)
add_bullet(doc, "Gunakan 3 sampai 5 KPI utama per peran agar fokus tetap jelas.")
add_bullet(doc, "Gabungkan hasil bisnis, kualitas proses, dan disiplin pencatatan. Hindari menilai hanya dari aktivitas.")
add_bullet(doc, "Tentukan periode, target, bobot, sumber data, pengecualian, dan pihak yang menyetujui koreksi.")
add_bullet(doc, "Jangan memakai target 100 persen untuk metrik yang secara operasional tidak sepenuhnya dapat dikendalikan individu.")

kpi_rows = [
    ("Digital marketing", "Qualified leads", "Jumlah leads yang memenuhi kriteria minimum", "CRM dan sumber kampanye", "30%"),
    ("Digital marketing", "Qualified lead rate", "Qualified leads dibagi total leads", "CRM", "25%"),
    ("Digital marketing", "Lead to opportunity", "Opportunity dari leads per sumber", "CRM dan pipeline", "25%"),
    ("Digital marketing", "Kelengkapan atribusi", "Leads dengan sumber dan kampanye lengkap", "CRM", "20%"),
    ("Sales online", "First response SLA", "Persentase chat pertama yang dijawab dalam batas waktu", "Inbox", "20%"),
    ("Sales online", "Follow-up compliance", "Lead aktif dengan next action tepat waktu", "CRM dan tugas", "20%"),
    ("Sales online", "Conversion rate", "Deal won dibagi opportunity qualified", "Pipeline", "25%"),
    ("Sales online", "Revenue atau gross profit", "Realisasi terhadap target yang disepakati", "Order", "35%"),
    ("Sales canvassing", "Kunjungan berkualitas", "Kunjungan dengan hasil, kebutuhan, dan tindak lanjut", "Aktivitas lapangan", "20%"),
    ("Sales canvassing", "Qualified opportunities", "Prospek yang lolos kriteria", "CRM", "20%"),
    ("Sales canvassing", "Conversion rate", "Deal won dibagi opportunity qualified", "Pipeline", "25%"),
    ("Sales canvassing", "Revenue atau gross profit", "Realisasi terhadap target", "Order", "35%"),
    ("PIC brand", "Requirement completeness", "Order dengan checklist lengkap sebelum proses", "Order dan checklist", "20%"),
    ("PIC brand", "First pass approval", "Sample disetujui tanpa revisi karena kesalahan internal", "Approval sample", "25%"),
    ("PIC brand", "Internal error rate", "Order dengan internal error dibagi order ditangani", "Error log", "35%"),
    ("PIC brand", "On-time handoff", "Handover ke produksi sesuai jadwal", "Order", "20%"),
    ("Produksi", "Schedule adherence", "Order atau batch selesai sesuai jadwal", "Produksi", "30%"),
    ("Produksi", "Output attainment", "Output aktual dibagi target yang disetujui", "Produksi", "25%"),
    ("Produksi", "Defect and rework rate", "Unit defect atau rework dibagi unit diperiksa", "QC", "30%"),
    ("Produksi", "Update discipline", "Progres dan kendala diperbarui tepat waktu", "Aktivitas", "15%"),
    ("Admin", "Unassigned lead rate", "Lead tanpa owner melewati batas waktu", "Inbox dan CRM", "25%"),
    ("Admin", "Data completeness", "Record wajib yang terisi lengkap", "CRM dan order", "25%"),
    ("Admin", "Overdue control", "Follow-up overdue yang ditindaklanjuti atau dieskalasi", "Tugas", "25%"),
    ("Admin", "Order document accuracy", "Dokumen order tanpa koreksi internal", "Order", "25%"),
]
add_table(doc, ["Peran", "KPI", "Definisi ringkas", "Sumber", "Bobot awal"], kpi_rows,
          [1.15, 1.35, 2.55, 1.25, 0.7], font_size=7.8,
          alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

add_heading(doc, "Rumus Skor KPI", 2)
add_paragraph(doc, "Untuk KPI yang semakin tinggi semakin baik, skor dapat dihitung sebagai realisasi dibagi target, lalu dibatasi sesuai kebijakan owner. Untuk KPI yang semakin rendah semakin baik, seperti defect rate, gunakan formula yang memberi nilai penuh ketika realisasi berada di bawah batas target dan menurun ketika batas terlampaui. Skor akhir adalah jumlah skor setiap KPI dikalikan bobotnya.")
add_paragraph(doc, "Insentif tidak disarankan dihitung dari revenue saja. Gunakan syarat kualitas minimum, misalnya tidak ada pelanggaran proses berat, data lengkap, dan tidak ada transaksi yang kemudian dibatalkan atau diretur karena kesalahan internal.")

add_heading(doc, "Rancangan OKR Awal", 1)
okr_rows = [
    ("Menjamin setiap lead mendapat penanganan yang jelas", "KR1: minimal persentase leads memiliki owner sesuai SLA. KR2: minimal persentase leads aktif memiliki next action. KR3: overdue turun dari baseline yang disepakati."),
    ("Meningkatkan efektivitas konversi penjualan", "KR1: qualified to won conversion mencapai target. KR2: alasan lost terisi lengkap. KR3: forecast dan realisasi berada dalam toleransi yang disepakati."),
    ("Mengurangi kesalahan order dan produksi", "KR1: seluruh order memiliki checklist requirement dan approval. KR2: internal error rate berada di bawah target. KR3: rework dan keterlambatan turun dari baseline."),
    ("Membuat kinerja tim dapat dipantau dari data", "KR1: KPI seluruh peran memiliki formula dan sumber data. KR2: dashboard diperbarui otomatis. KR3: review mingguan menggunakan data yang sama."),
]
add_table(doc, ["Objective", "Contoh Key Results"], okr_rows, [2.35, 4.45], font_size=9.0)
add_paragraph(doc, "Angka final untuk setiap key result harus disetujui setelah baseline tersedia. KPI digunakan untuk evaluasi kinerja rutin, sedangkan OKR digunakan untuk mendorong perbaikan penting dalam satu periode.", style="Small Note")

add_heading(doc, "Data Minimum yang Harus Disimpan", 1)
data_rows = [
    ("Kontak", "Nama, nomor atau akun sosial, perusahaan atau brand, sumber, izin komunikasi, pemilik"),
    ("Lead", "Kebutuhan, produk atau layanan, volume, anggaran, deadline, kualitas lead, status, next action"),
    ("Percakapan", "Kanal, waktu, pengirim, penerima, isi atau referensi pesan, assignment, waktu respons"),
    ("Opportunity", "Nilai, kuantitas, produk, tahap, probabilitas, target close, quotation, alasan lost"),
    ("Order", "PO, item, spesifikasi, harga, jumlah, tenggat, rute trading atau produksi, status pembayaran"),
    ("Produksi", "Vendor, material, target, output, jadwal, kendala, QC, rework, bukti progres"),
    ("KPI", "Periode, peran, indikator, formula, target, realisasi, bobot, skor, sumber, approval"),
]
add_table(doc, ["Objek", "Field minimum"], data_rows, [1.3, 5.5], font_size=9.1)

add_heading(doc, "Kebutuhan Nonfungsional", 1)
add_table(doc,
          ["Area", "Kebutuhan awal"],
          [
              ("Keamanan", "Login individual, hak akses berbasis peran, pembatasan data sensitif, dan pencabutan akses ketika personel keluar."),
              ("Audit", "Riwayat perubahan status, owner, nilai, target, formula, dan approval tidak dapat dihapus oleh pengguna biasa."),
              ("Ketersediaan", "Dashboard dapat digunakan pada jam operasional dengan mekanisme antrean ketika kanal eksternal mengalami gangguan."),
              ("Kinerja", "Inbox dan daftar leads tetap responsif pada volume 30 sampai 100 leads per hari serta pertumbuhan historis."),
              ("Privasi", "Nomor telepon dan isi percakapan hanya terlihat oleh peran yang berwenang; ekspor dicatat."),
              ("Integrasi", "Koneksi kanal memakai akses resmi, akun bisnis yang memenuhi syarat, dan persetujuan yang diperlukan dari penyedia kanal."),
              ("Backup", "Data operasional dicadangkan dan memiliki prosedur pemulihan yang diuji."),
              ("Kemudahan pakai", "Status, istilah, warna, dan tindakan konsisten; layar utama menampilkan pekerjaan berikutnya, bukan hanya statistik."),
          ], [1.35, 5.45], font_size=9.1)

add_heading(doc, "Notifikasi dan Eskalasi", 2)
add_bullet(doc, "Lead baru belum memiliki owner melewati SLA.")
add_bullet(doc, "Pesan pelanggan belum dijawab atau follow-up melewati tenggat.")
add_bullet(doc, "Hot lead tidak memiliki aktivitas dalam batas waktu.")
add_bullet(doc, "Requirement order belum lengkap menjelang jadwal produksi.")
add_bullet(doc, "Produksi tertinggal dari target atau memiliki kendala terbuka.")
add_bullet(doc, "QC menemukan defect atau rework di atas batas.")
add_bullet(doc, "Tanggal kirim berisiko atau terlewati.")

add_heading(doc, "Tahapan Implementasi", 1)
roadmap = [
    ("Fase 0 Validasi", "Konfirmasi proses, struktur tim, definisi status, hak akses, baseline, KPI, integrasi kanal, dan data lama.", "Dokumen kebutuhan disetujui"),
    ("Fase 1 Leads", "Inbox omnichannel, kontak, leads, assignment, follow-up, pipeline, dan dashboard owner dasar.", "Tidak ada lead tanpa owner dan tindak lanjut"),
    ("Fase 2 Order", "Order, requirement, approval, kolaborasi sales, trading, dan handover ke operasional.", "Deal won dapat ditelusuri menjadi order"),
    ("Fase 3 Produksi", "Jadwal, vendor konveksi, material, progres, QC, rework, pengiriman, dan error log.", "Order dapat dipantau sampai selesai"),
    ("Fase 4 Kinerja", "Target, KPI, OKR, skor, insentif, laporan tren, dan evaluasi berkala.", "Skor berasal dari data operasional"),
]
add_table(doc, ["Fase", "Cakupan", "Kriteria selesai"], roadmap, [1.25, 3.75, 1.8], font_size=8.9)

add_heading(doc, "Prioritas MVP", 2)
add_paragraph(doc, "MVP disarankan berisi inbox omnichannel, kontak dan leads, assignment, follow-up, pipeline, dashboard owner dasar, hak akses, dan audit log. Modul produksi dapat dimulai dengan pencatatan status sederhana setelah proses order disepakati. KPI otomatis sebaiknya dirilis setelah definisi data dan formula stabil.")

add_heading(doc, "Kriteria Penerimaan Awal", 1)
criteria = [
    "Admin dapat melihat percakapan dari kanal yang disepakati pada satu inbox sesuai hak akses.",
    "Setiap pesan baru dapat dihubungkan ke kontak dan lead tanpa kehilangan asal kanal.",
    "Setiap lead memiliki owner, status, next action, dan riwayat aktivitas yang dapat ditelusuri.",
    "Sistem menampilkan leads yang belum ditangani, belum dijawab, atau overdue.",
    "Sales dapat meminta kolaborasi dan mencatat kontribusi tanpa membuat peluang duplikat.",
    "Deal won dapat diubah menjadi order dengan data requirement dan approval.",
    "Owner dapat melihat target dan realisasi sales, order berisiko, progres produksi, dan KPI pada periode yang dipilih.",
    "Perubahan data penting tersimpan di audit log dengan pengguna dan waktu perubahan.",
    "Laporan dapat difilter berdasarkan periode, kanal, sales, layanan, status, dan sumber lead.",
]
for item in criteria:
    add_bullet(doc, item)

add_heading(doc, "Risiko dan Mitigasi", 1)
add_table(doc,
          ["Risiko", "Dampak", "Mitigasi"],
          [
              ("Definisi proses belum seragam", "Dashboard menampilkan angka yang diperdebatkan", "Setujui kamus status, formula, dan owner proses sebelum pembangunan"),
              ("Data lama tidak konsisten", "Duplikasi dan laporan awal tidak akurat", "Mapping, cleansing, deduplikasi, dan uji sampel sebelum migrasi"),
              ("Ketergantungan kanal eksternal", "Pesan terlambat atau fitur terbatas", "Gunakan koneksi resmi, antrean, retry, monitoring, dan prosedur manual darurat"),
              ("KPI mendorong perilaku yang salah", "Tim mengejar aktivitas atau revenue dengan mengorbankan kualitas", "Gunakan metrik berimbang dan review berkala"),
              ("Tim tidak disiplin memperbarui data", "Dashboard tidak dapat dipercaya", "Sederhanakan input, wajibkan field inti, otomatisasi, dan audit kepatuhan"),
              ("Akses chat terlalu luas", "Risiko privasi dan penyalahgunaan data", "Role-based access, masking, audit ekspor, dan review akses"),
          ], [2.35, 1.95, 2.5], font_size=8.35)

add_heading(doc, "Keputusan yang Diperlukan dari Owner", 1)
decisions = [
    ("Struktur tim", "Nama, jumlah orang, atasan, peran rangkap, dan wilayah kerja setiap personel."),
    ("Akses chat", "Siapa dapat melihat seluruh chat, chat tim, chat sendiri, data kontak, dan ekspor."),
    ("SLA", "Batas waktu assignment, respons pertama, follow-up, dan eskalasi per jam operasional."),
    ("Kriteria lead", "Definisi new, qualified, warm, hot, lost, cancelled, dan dormant."),
    ("Kolaborasi sales", "Aturan primary owner, contributor, handover, komisi, dan penyelesaian konflik."),
    ("Zero mistakes", "Daftar kesalahan, toleransi, sumber penyebab, pengecualian, dan dampak skor."),
    ("Target dan insentif", "Target per peran, periode, bobot KPI, batas skor, formula bonus, dan approval."),
    ("Proses produksi", "Tahap wajib, PIC, vendor, jenis QC, bukti progres, serta definisi on time."),
    ("Integrasi", "Akun bisnis, nomor, halaman, profil, administrator, dan kesiapan akses kanal."),
    ("Migrasi data", "Spreadsheet yang menjadi sumber utama, periode data, dan aturan pembersihan."),
]
add_table(doc, ["Keputusan", "Hal yang perlu ditetapkan"], decisions, [1.6, 5.2], font_size=9.2)

add_heading(doc, "Pertanyaan Validasi Berikutnya", 1)
questions = [
    "Apakah kebutuhan sebenarnya adalah admin dapat memantau semua chat? Siapa lagi yang memiliki hak yang sama?",
    "Berapa jumlah personel saat ini dan siapa yang menjalankan digital marketing, sales online, canvassing, PIC brand, produksi, dan admin?",
    "Apakah 30 sampai 100 leads merupakan rata-rata harian, puncak harian, atau target? Berapa proporsi setiap kanal?",
    "Apa kriteria minimum agar lead dinilai qualified, warm, atau hot?",
    "Apa target respons pertama dan follow-up pada jam kerja serta di luar jam kerja?",
    "Bagaimana komisi dibagi ketika sales online dan canvassing berkolaborasi?",
    "Apa jenis kesalahan PIC brand yang termasuk internal error, vendor error, atau perubahan pelanggan?",
    "Apakah target utama sales menggunakan revenue, gross profit, kuantitas, atau kombinasi?",
    "Bagaimana alur pembayaran, termin, pembelian kain, dan pencatatan biaya pada order?",
    "Spreadsheet mana yang menjadi sumber data utama dan data historis berapa lama yang perlu dimigrasikan?",
]
for q in questions:
    add_bullet(doc, q)

add_heading(doc, "Langkah Berikutnya", 1)
reset_numbering()
add_number(doc, "Owner dan tim memvalidasi interpretasi, cakupan, peran, status, serta daftar keputusan dalam dokumen ini.")
add_number(doc, "Lakukan workshop proses singkat untuk menutup pertanyaan validasi dan menyepakati kamus data.")
add_number(doc, "Buat prototype layar utama untuk inbox, detail lead, pipeline, dashboard owner, dan detail order.")
add_number(doc, "Uji prototype dengan contoh kasus nyata dari chat masuk sampai order selesai.")
add_number(doc, "Tetapkan backlog MVP, estimasi, integrasi, rencana migrasi, dan kriteria penerimaan final.")

add_heading(doc, "Kesimpulan", 1)
add_paragraph(doc, "Silancar tidak hanya membutuhkan penggabungan chat, tetapi juga satu alur data dari pesan masuk sampai hasil penjualan dan pemenuhan order. Prioritas awal adalah memastikan setiap lead memiliki owner, tindak lanjut, dan status yang jelas. Setelah fondasi tersebut stabil, data order, produksi, KPI, dan OKR dapat dihitung secara konsisten dan dipakai owner untuk mengambil keputusan.")

# Footer with document title and dynamic page field.
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Silancar  Hasil Interview dan Rancangan Awal Solusi   |   ")
    set_font(r, size=8, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r2 = p.add_run()
    r2._r.append(fld_char1)
    r2._r.append(instr_text)
    r2._r.append(fld_char2)
    set_font(r2, size=8, color=MID_GRAY)

# Core properties and save.
doc.core_properties.title = "Silancar Hasil Interview dan Rancangan Awal Solusi"
doc.core_properties.subject = "Rancangan omnichannel, sales, operasional, produksi, KPI, dan OKR"
doc.core_properties.author = ""
doc.core_properties.keywords = "Silancar, omnichannel, sales, produksi, KPI, OKR"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT.resolve())
